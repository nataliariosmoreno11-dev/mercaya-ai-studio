import io, zipfile, re, unicodedata, hashlib, math
from pathlib import Path
import pandas as pd
from docx import Document

REQUIRED=['01_business_hourly.csv','02_paid_media.csv','03_regions.csv','04_coupons.csv','05_crm.csv','06_competitors.csv']

def norm(s):
    return ''.join(c for c in unicodedata.normalize('NFKD',str(s)) if not unicodedata.combining(c)).strip().lower()

def read_sources(upload):
    data={}; brief=''; names=[]
    if upload:
        with zipfile.ZipFile(upload) as z:
            for name in z.namelist():
                short=name.split('/')[-1]
                if short in REQUIRED: data[short]=pd.read_csv(io.BytesIO(z.read(name)));names.append(short)
                if short=='07_campaign_brief.docx': brief='\n'.join(p.text for p in Document(io.BytesIO(z.read(name))).paragraphs)
    else:
        from pathlib import Path
        for short in REQUIRED:
            p=Path(__file__).parent/'data'/short
            if p.exists():data[short]=pd.read_csv(p);names.append(short)
        p=Path(__file__).parent/'data'/'07_campaign_brief.docx'
        if p.exists():brief='\n'.join(p.text for p in Document(p).paragraphs)
    return data,brief,names

def inspect(data,brief):
    alerts=[]
    def finding(name,priority,fact,hypothesis,missing,action,owner,confidence='Alta',category='Negocio'):
        alerts.append(dict(Senal=name,Prioridad=priority,Hecho=fact,Hipotesis=hypothesis,Dato_faltante=missing,Accion=action,Responsable=owner,Confianza=confidence,Categoria=category,Estado='Pendiente de validación'))
    paid=data.get('02_paid_media.csv');biz=data.get('01_business_hourly.csv');regions=data.get('03_regions.csv');coupons=data.get('04_coupons.csv');crm=data.get('05_crm.csv')
    if paid is not None:
        duplicates=int(paid.duplicated().sum())
        if duplicates:finding('Filas duplicadas en medios','Alta',f'{duplicates} fila(s) idénticas duplicadas en Paid Media.','La duplicación puede inflar inversión y atribución.','Confirmar clave de unicidad de la extracción.','Excluir duplicados exactos del análisis; pedir validación al área de datos.','Data / Paid Media',category='Calidad')
        paid=paid.drop_duplicates().copy()
    if regions is not None:
        aliases=[(v,norm(v)) for v in regions['region'].astype(str)]
        if len(set(x[0] for x in aliases))>len(set(x[1] for x in aliases)):
            finding('Nombres inconsistentes de regiones','Media','Hay regiones con etiquetas distintas pero equivalentes al normalizar tildes.','Puede fragmentar reportes por región.','Confirmar si las filas son segmentos distintos o registros de una misma región.','Homologar etiquetas, sin sumar registros potencialmente superpuestos hasta validar grano.','Data / Comercial',category='Calidad')
    if coupons is not None:
        flash=coupons[coupons['coupon'].astype(str).str.contains('FLASH1010',case=False,na=False)]
        if len(flash):
            system=str(flash.iloc[0]['system_validity_window']); crm_text=' '.join(crm.loc[crm['message'].astype(str).str.contains('FLASH1010',case=False,na=False),'message'].astype(str)) if crm is not None else ''
            crm_window=re.search(r'\b\d{2}:\d{2}\s*[-–]\s*\d{2}:\d{2}\b',crm_text)
            brief_window=re.search(r'(?:Ventana comunicada en brief:|FLASH1010.{0,100}?)[^\n]{0,120}?(\d{2}:\d{2}\s*[-–]\s*\d{2}:\d{2})',brief,re.I)
            b=brief_window.group(1) if brief_window else 'no extraída'; c=crm_window.group(0) if crm_window else 'no identificada'
            if re.sub(r'\s|–','',norm(b))!=re.sub(r'\s|–','',norm(system)) or (crm_window and re.sub(r'\s|–','',norm(c))!=re.sub(r'\s|–','',norm(b))):
                finding('Ventanas FLASH1010 contradictorias','Crítica',f'Brief: {b}; sistema: {system}; comunicación CRM: {c}.','Una comunicación adicional podría anunciar horarios no válidos.','Confirmación de la ventana oficial y T&C vigentes.','Generar incidencia y solicitar autorización para pausar nuevas comunicaciones; validar términos oficiales.','CRM / Promociones / Marketing Lead',category='Promociones')
            consumed=float(flash.iloc[0]['consumed_cop_mm']); budget=float(flash.iloc[0]['budget_cop_mm']);rate=consumed/budget if budget else 0
            if rate>=.8:finding('Presupuesto FLASH1010 comprometido','Alta',f'Se consumieron {consumed:g} de {budget:g} MM COP ({rate:.0%}).','La demanda podría agotar el cupón antes del cierre.','Ritmo de redenciones por minuto y demanda futura.','Monitorear consumo y solicitar revisión de Finanzas antes de aumentar presupuesto.','Promociones / Finanzas',category='Promociones')
    if biz is not None:
        b=biz.sort_values('timestamp'); a=float(b.iloc[0]['actual_cvr']);last=float(b.iloc[-1]['actual_cvr']);s=float(b['actual_sessions'].sum());o=float(b['orders'].sum());cvr=o/s if s else 0
        if last<a:finding('Deterioro de conversión','Alta',f'CVR reportada pasó de {a:.2%} a {last:.2%}; CVR ponderada calculada = {cvr:.2%} en horas observadas.','Podría haber fricción en checkout, oferta, medios o promociones.','Abandono por etapa, cohortes, errores y fuentes de tráfico.','Analizar embudo por franja y canal, validar fricción y probar correcciones controladas.','E-commerce / Producto / Marketing',confidence='Media')
        if regions is not None:
            total=float(regions['actual_nmv_cop_mm'].sum());hours=float(b['actual_nmv_cop_mm'].sum());
            if abs(total-hours)>.001:finding('Fuentes de NMV no conciliadas','Alta',f'NMV regional = {total:g} MM; NMV horario = {hours:g} MM; diferencia = {total-hours:+g} MM.','Las fuentes podrían tener distintos cortes o granularidades.','Ventana, filtros y definiciones de cada fuente.','No combinar totales; solicitar conciliación a Data antes de reportar NMV único.','Data / BI',category='Calidad')
    if paid is not None:
        sums=paid.groupby('channel',as_index=False)[['spend_cop_mm','attributed_revenue_cop_mm']].sum();sums['ROAS']=sums['attributed_revenue_cop_mm']/sums['spend_cop_mm'].replace(0,float('nan'))
        if sums['ROAS'].dropna().nunique()>1:finding('Diferencias de eficiencia entre canales','Media','ROAS agregado por canal muestra diferencias; ingresos son atribuidos y no equivalen a NMV incremental.','Puede haber una oportunidad de optimización del mix.','Incrementalidad, pacing, saturación y límites de presupuesto.','Revisar eficiencia e incrementalidad antes de proponer cambios; no mover presupuesto automáticamente.','Paid Media / Marketing Lead',confidence='Media')
    order={'Crítica':0,'Alta':1,'Media':2,'Baja':3};alerts.sort(key=lambda d:order.get(d['Prioridad'],4))
    return data,paid,alerts


SCHEMAS = {
 REQUIRED[0]: {'timestamp':'date','actual_sessions':'num','orders':'num','actual_cvr':'rate','actual_nmv_cop_mm':'num','forecast_nmv_cop_mm':'num','forecast_cvr':'rate'},
 REQUIRED[1]: {'channel':'text','spend_cop_mm':'num','attributed_revenue_cop_mm':'num','sessions':'num'},
 REQUIRED[2]: {'region':'text','actual_nmv_cop_mm':'num'},
 REQUIRED[3]: {'coupon':'text','budget_cop_mm':'num','consumed_cop_mm':'num','system_validity_window':'text'},
 REQUIRED[4]: {'message':'text'}, REQUIRED[5]: {'competitor':'text','promotion':'text','source_status':'text'}
}

def validate(data):
    valid={}; checks=[]
    for name, schema in SCHEMAS.items():
        df=data.get(name); errors=[]
        if df is None: errors=['Fuente ausente']
        else:
            df=df.copy()
            if df.empty: errors.append('Sin filas')
            for col,kind in schema.items():
                if col not in df: errors.append(f'Falta columna {col}'); continue
                if df[col].isna().any(): errors.append(f'Nulos en {col}')
                if kind in ('num','rate'):
                    v=pd.to_numeric(df[col],errors='coerce')
                    if (v.isna() | ~v.map(lambda x: math.isfinite(x)) | (v<0)).any(): errors.append(f'Número inválido en {col}')
                    if kind=='rate' and (v>1).any(): errors.append(f'Tasa fuera de [0,1] en {col}')
                    df[col]=v
                if kind=='date':
                    if pd.to_datetime(df[col],errors='coerce').isna().any(): errors.append(f'Fecha inválida en {col}')
                    elif pd.to_datetime(df[col]).duplicated().any(): errors.append('Timestamp duplicado; grano ambiguo')
                    else: df[col]=pd.to_datetime(df[col]).dt.strftime('%Y-%m-%d %H:%M:%S')
                if kind=='text' and df[col].astype(str).str.strip().eq('').any(): errors.append(f'Texto vacío en {col}')
            if not errors: valid[name]=df
        checks.append({'Fuente':name,'Estado':'BLOQUEADA' if errors else 'Válida','Detalle':'; '.join(errors) or f'{len(df)} filas; esquema y dominios válidos'})
    return valid,checks

EVIDENCE={
 'Filas duplicadas en medios':[REQUIRED[1]], 'Nombres inconsistentes de regiones':[REQUIRED[2]],
 'Ventanas FLASH1010 contradictorias':[REQUIRED[3],REQUIRED[4],'07_campaign_brief.docx'],
 'Presupuesto FLASH1010 comprometido':[REQUIRED[3]], 'Deterioro de conversión':[REQUIRED[0]],
 'Fuentes de NMV no conciliadas':[REQUIRED[0],REQUIRED[2]], 'Diferencias de eficiencia entre canales':[REQUIRED[1]]}
RULES={
 'Filas duplicadas en medios':'R01: duplicados exactos > 0; excluir del agregado de medios.',
 'Nombres inconsistentes de regiones':'R02: etiquetas distintas con igual normalización; no fusionar.',
 'Ventanas FLASH1010 contradictorias':'R03: ventanas extraídas distintas o no identificadas; requiere validar términos.',
 'Presupuesto FLASH1010 comprometido':'R04: consumo / presupuesto >= 80%.',
 'Deterioro de conversión':'R05: CVR última hora < primera hora; no demuestra causalidad.',
 'Fuentes de NMV no conciliadas':'R06: diferencia absoluta entre sumas > 0.001 MM COP; cortes no confirmados.',
 'Diferencias de eficiencia entre canales':'R07: ROAS agregado distinto en al menos dos canales con inversión positiva; revisión exploratoria, no incrementalidad.'}

def analyze(data,brief):
    valid,checks=validate(data)
    _,paid,alerts=inspect(valid,brief)
    for a in alerts:
        a['Regla']=RULES[a['Senal']];a['Fuentes']=EVIDENCE[a['Senal']]
        a['ID']=a['Regla'].split(':')[0]
        a['Aprobacion']='Pendiente: '+('Finanzas y Marketing Lead' if a['Categoria']=='Promociones' else a['Responsable'])
        a['Justificacion']='Prioridad fija de la regla: '+a['Prioridad']+'. '+a['Hecho']
    return valid,paid,alerts,checks

def fingerprint(data,brief):
    body=''.join(k+v.to_csv(index=False) for k,v in sorted(data.items()))+brief
    return hashlib.sha256(body.encode()).hexdigest()

def artifacts(alerts,source_id):
    lines=['MercaYa Intelligence | Brief de ejecución','BORRADOR — requiere aprobación humana',f'Paquete SHA256: {source_id}',
           'Motor determinístico; sin IA generativa en ejecución. No se ejecutan cambios comerciales.','']
    rows=[]
    for a in alerts:
        for key in ['ID','Senal','Prioridad','Hecho','Hipotesis','Dato_faltante','Accion','Responsable','Regla','Justificacion','Aprobacion']:
            lines.append(f'{key}: {a[key]}')
        lines += ['Evidencia: '+', '.join(a['Fuentes']), 'Estado: Borrador pendiente de revisión','']
        rows.append({**a,'Fuentes':', '.join(a['Fuentes']),'Paquete_SHA256':source_id,'Estado':'Borrador pendiente de revisión','Aprobado_por':'','Fecha_aprobacion':'','Resultado':''})
    return '\n'.join(lines),pd.DataFrame(rows).to_csv(index=False).encode('utf-8-sig')
